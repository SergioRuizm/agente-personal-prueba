"""
Agente personal (versión Google Gemini): organización del día a día + estudio.

Instalación (en la terminal de VS Code):
    pip install langchain-google-genai langgraph

Clave de API (gratis en aistudio.google.com), en PowerShell:
    $env:GOOGLE_API_KEY="tu_clave"

Uso:
    python agente_personal_gemini.py

Guarda sus datos en la carpeta "datos_agente" (tareas.json, apuntes.json, tarjetas.json).
"""

import difflib
import json
import re
from datetime import datetime
from pathlib import Path

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.tools import tool
from langgraph.prebuilt import create_react_agent

# Librerías para leer documentos (pip install pypdf python-docx)
try:
    import pypdf  # PDF
except ImportError:
    pypdf = None
try:
    import docx  # Word (.docx)
except ImportError:
    docx = None
try:
    from PIL import Image, ImageColor, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps  # fotos
except ImportError:
    Image = None

# ---------- Almacenamiento sencillo en archivos JSON ----------

CARPETA = Path("datos_agente")
CARPETA.mkdir(exist_ok=True)


def cargar(nombre: str) -> list:
    ruta = CARPETA / f"{nombre}.json"
    if ruta.exists():
        return json.loads(ruta.read_text(encoding="utf-8"))
    return []


def guardar(nombre: str, datos: list) -> None:
    ruta = CARPETA / f"{nombre}.json"
    ruta.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------- Herramientas: día a día y organización ----------

@tool
def fecha_hora_actual() -> str:
    """Devuelve la fecha y hora actuales. Úsala siempre que necesites saber qué día es hoy."""
    return datetime.now().strftime("%A %d/%m/%Y, %H:%M")


@tool
def anadir_tarea(titulo: str, fecha_limite: str = "", categoria: str = "general") -> str:
    """Añade una tarea pendiente. fecha_limite en formato AAAA-MM-DD (opcional).
    categoria puede ser 'estudio', 'personal', 'recados', etc."""
    tareas = cargar("tareas")
    nuevo_id = max([t["id"] for t in tareas], default=0) + 1
    tareas.append({
        "id": nuevo_id,
        "titulo": titulo,
        "fecha_limite": fecha_limite,
        "categoria": categoria,
        "hecha": False,
    })
    guardar("tareas", tareas)
    return f"Tarea #{nuevo_id} añadida: {titulo}"


@tool
def listar_tareas(solo_pendientes: bool = True) -> str:
    """Lista las tareas guardadas, ordenadas por fecha límite.
    Con solo_pendientes=True oculta las ya completadas."""
    tareas = cargar("tareas")
    if solo_pendientes:
        tareas = [t for t in tareas if not t["hecha"]]
    if not tareas:
        return "No hay tareas."
    tareas.sort(key=lambda t: t["fecha_limite"] or "9999-99-99")
    return "\n".join(
        f"#{t['id']} [{'x' if t['hecha'] else ' '}] {t['titulo']} "
        f"(límite: {t['fecha_limite'] or 'sin fecha'}, {t['categoria']})"
        for t in tareas
    )


@tool
def completar_tarea(id_tarea: int) -> str:
    """Marca una tarea como completada usando su número (id)."""
    tareas = cargar("tareas")
    for t in tareas:
        if t["id"] == id_tarea:
            t["hecha"] = True
            guardar("tareas", tareas)
            return f"Tarea #{id_tarea} completada: {t['titulo']}"
    return f"No existe la tarea #{id_tarea}."


# ---------- Herramientas: estudio ----------

@tool
def guardar_apunte(asignatura: str, texto: str) -> str:
    """Guarda un apunte o resumen de estudio asociado a una asignatura."""
    apuntes = cargar("apuntes")
    apuntes.append({
        "asignatura": asignatura,
        "fecha": datetime.now().strftime("%Y-%m-%d"),
        "texto": texto,
    })
    guardar("apuntes", apuntes)
    return f"Apunte guardado en {asignatura}."


@tool
def leer_apuntes(asignatura: str = "") -> str:
    """Devuelve los apuntes guardados. Si se indica asignatura, solo los de esa asignatura."""
    apuntes = cargar("apuntes")
    # El número de cada apunte es su posición en la lista completa (empieza en 1)
    numerados = list(enumerate(apuntes, start=1))
    if asignatura:
        numerados = [(n, a) for n, a in numerados if a["asignatura"].lower() == asignatura.lower()]
    if not numerados:
        return "No hay apuntes."
    return "\n\n".join(
        f"#{n} [{a['asignatura']} · {a['fecha']}]\n{a['texto']}" for n, a in numerados
    )


@tool
def eliminar_apunte(numero: int) -> str:
    """Elimina UN apunte usando su número (el que aparece como #número al leer los apuntes).
    Antes de llamar a esta herramienta, usa leer_apuntes para ver los números.
    El programa pedirá confirmación al usuario por teclado antes de borrar."""
    apuntes = cargar("apuntes")
    if numero < 1 or numero > len(apuntes):
        return f"No existe el apunte #{numero}."
    apunte = apuntes[numero - 1]
    vista_previa = apunte["texto"][:150]
    print(f"\n[Confirmación] Se va a eliminar el apunte #{numero} de {apunte['asignatura']}:")
    print(f"  {vista_previa}{'...' if len(apunte['texto']) > 150 else ''}")
    respuesta = input("¿Seguro que quieres eliminarlo? (s/n): ").strip().lower()
    if respuesta not in {"s", "si", "sí"}:
        return "Cancelado: el usuario no confirmó, no se ha borrado nada."
    apuntes.pop(numero - 1)
    guardar("apuntes", apuntes)
    return f"Apunte #{numero} eliminado. Los números de los apuntes siguientes han cambiado."


@tool
def guardar_tarjetas(asignatura: str, preguntas: list[str], respuestas: list[str]) -> str:
    """Guarda tarjetas de repaso. 'preguntas' y 'respuestas' son dos listas de texto
    con el MISMO número de elementos: la respuesta nº 1 corresponde a la pregunta nº 1.
    Primero genera las tarjetas a partir de los apuntes y luego llama a esta herramienta."""
    if len(preguntas) != len(respuestas):
        return "Error: debe haber el mismo número de preguntas y respuestas."
    guardadas = cargar("tarjetas")
    for p, r in zip(preguntas, respuestas):
        guardadas.append({"asignatura": asignatura, "pregunta": p, "respuesta": r})
    guardar("tarjetas", guardadas)
    return f"{len(preguntas)} tarjetas guardadas en {asignatura}."


@tool
def leer_tarjetas(asignatura: str = "") -> str:
    """Devuelve las tarjetas de repaso guardadas, para hacer un test o repaso al usuario."""
    tarjetas = cargar("tarjetas")
    if asignatura:
        tarjetas = [t for t in tarjetas if t["asignatura"].lower() == asignatura.lower()]
    if not tarjetas:
        return "No hay tarjetas."
    return "\n".join(f"P: {t['pregunta']} | R: {t['respuesta']}" for t in tarjetas)


# ---------- Herramientas: revisar código ----------

CARPETA_CODIGO = Path("codigo")  # el usuario deja aquí los archivos que quiere que revise
CARPETA_CODIGO.mkdir(exist_ok=True)
EXTENSIONES_CODIGO = {".py", ".js", ".ts", ".html", ".css", ".java", ".c", ".cpp", ".cs",
                      ".php", ".sql", ".json", ".txt", ".md"}


@tool
def listar_archivos_codigo() -> str:
    """Lista los archivos de código que el usuario ha dejado en la carpeta 'codigo'
    para que los revises. Úsala cuando pida revisar o corregir código sin dar el nombre."""
    archivos = sorted(
        f.name for f in CARPETA_CODIGO.iterdir()
        if f.is_file() and f.suffix.lower() in EXTENSIONES_CODIGO
    )
    if not archivos:
        return "La carpeta 'codigo' está vacía. El usuario debe copiar ahí el archivo a revisar."
    return "\n".join(archivos)


@tool
def leer_archivo_codigo(nombre: str) -> str:
    """Lee un archivo de código de la carpeta 'codigo' (solo el nombre, por ejemplo 'ejercicio1.py')
    y lo devuelve con números de línea, para poder revisarlo y señalar errores concretos.
    Solo lee: NO modifica el archivo."""
    ruta = CARPETA_CODIGO / Path(nombre).name  # .name evita salir de la carpeta
    if not ruta.is_file() or ruta.suffix.lower() not in EXTENSIONES_CODIGO:
        return f"No encuentro '{nombre}' en la carpeta 'codigo'."
    texto = ruta.read_text(encoding="utf-8", errors="replace")
    if len(texto) > 20000:
        return "El archivo es demasiado largo (más de 20.000 caracteres). Pide al usuario una parte."
    lineas = texto.splitlines()
    return "\n".join(f"{i}: {linea}" for i, linea in enumerate(lineas, start=1))


@tool
def guardar_codigo_corregido(nombre: str, codigo_corregido: str) -> str:
    """Aplica una corrección a un archivo que YA existe en la carpeta 'codigo'.
    'codigo_corregido' debe ser el contenido COMPLETO del archivo ya corregido (no solo las
    partes que cambian) y SIN números de línea. Úsala solo cuando el usuario pida corregir o
    arreglar el archivo, después de haberlo leído con leer_archivo_codigo. Antes de guardar,
    el programa enseña los cambios al usuario, le pide confirmación y guarda una copia de
    seguridad del original."""
    ruta = CARPETA_CODIGO / Path(nombre).name
    if not ruta.is_file() or ruta.suffix.lower() not in EXTENSIONES_CODIGO:
        return f"No encuentro '{nombre}' en la carpeta 'codigo'."

    nuevas = codigo_corregido.splitlines()
    con_numero = sum(1 for linea in nuevas if re.match(r"^\d+: ", linea))
    if nuevas and con_numero > len(nuevas) / 2:
        return "Error: el código lleva números de línea ('12: ...'). Envíalo sin ellos."

    original = ruta.read_text(encoding="utf-8", errors="replace")
    if codigo_corregido.strip() == original.strip():
        return "El archivo ya está igual: no hay cambios que aplicar."

    diferencias = list(difflib.unified_diff(
        original.splitlines(), nuevas,
        fromfile=f"{ruta.name} (original)", tofile=f"{ruta.name} (corregido)", lineterm="",
    ))
    print(f"\n[Confirmación] Cambios propuestos en {ruta.name}:")
    for linea in diferencias[:80]:
        print(linea)
    if len(diferencias) > 80:
        print("... (hay más cambios)")
    if len(codigo_corregido) < len(original) * 0.5:
        print("\nOJO: el código corregido es mucho más corto que el original. "
              "Puede que falten partes.")
    respuesta = input("\n¿Aplicar estos cambios? (s/n): ").strip().lower()
    if respuesta not in {"s", "si", "sí"}:
        return "Cancelado: el usuario no confirmó, el archivo no se ha modificado."

    copias = CARPETA_CODIGO / "copias_seguridad"
    copias.mkdir(exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    copia = copias / f"{ruta.stem}_{marca}{ruta.suffix}"
    copia.write_text(original, encoding="utf-8")
    ruta.write_text(codigo_corregido, encoding="utf-8")
    return f"Archivo {ruta.name} corregido. Copia del original guardada en {copia}."


# ---------- Herramientas: documentos subidos ----------

CARPETA_DOCS = Path("documentos")  # aquí se guardan los archivos que sube el usuario
CARPETA_DOCS.mkdir(exist_ok=True)
EXTENSIONES_DOCS = {".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".json", ".xml"}
TAMANO_PARTE = 12000  # caracteres que se leen de una vez


def extraer_texto(ruta: Path) -> str:
    """Saca el texto de un PDF, un Word o un archivo de texto."""
    sufijo = ruta.suffix.lower()
    if sufijo == ".pdf":
        if pypdf is None:
            return "ERROR: falta la librería pypdf (pip install pypdf)."
        lector = pypdf.PdfReader(str(ruta))
        paginas = []
        for n, pagina in enumerate(lector.pages, start=1):
            paginas.append(f"--- Página {n} ---\n{pagina.extract_text() or ''}")
        return "\n".join(paginas)
    if sufijo == ".docx":
        if docx is None:
            return "ERROR: falta la librería python-docx (pip install python-docx)."
        documento = docx.Document(str(ruta))
        partes = [p.text for p in documento.paragraphs]
        for tabla in documento.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(celda.text for celda in fila.cells))
        return "\n".join(partes)
    return ruta.read_text(encoding="utf-8", errors="replace")


@tool
def listar_documentos() -> str:
    """Lista los documentos que el usuario ha subido (carpeta 'documentos'):
    PDF, Word (.docx), texto, etc. Úsala cuando mencione un archivo o documento sin dar el nombre exacto."""
    archivos = sorted(
        f for f in CARPETA_DOCS.iterdir()
        if f.is_file() and f.suffix.lower() in EXTENSIONES_DOCS
    )
    if not archivos:
        return "No hay documentos subidos."
    return "\n".join(f"{f.name} ({f.stat().st_size // 1024} KB)" for f in archivos)


@tool
def leer_documento(nombre: str, parte: int = 1) -> str:
    """Lee el texto de un documento subido (PDF, Word .docx o texto). Si es largo, se lee por
    partes: pide parte=1, parte=2... hasta el final antes de resumirlo. 'nombre' es el nombre
    del archivo, por ejemplo 'bases_concurso.pdf'. Solo lee: no modifica el archivo."""
    ruta = CARPETA_DOCS / Path(nombre).name  # .name evita salir de la carpeta
    if not ruta.is_file() or ruta.suffix.lower() not in EXTENSIONES_DOCS:
        return f"No encuentro '{nombre}' entre los documentos subidos."
    try:
        texto = extraer_texto(ruta)
    except Exception as error:
        return f"No se pudo leer '{nombre}': {error}"
    if texto.startswith("ERROR:"):
        return texto
    if len(texto.strip()) < 20 or (ruta.suffix.lower() == ".pdf" and len(texto.replace("--- Página", "").strip()) < 40):
        return ("El documento no tiene texto que pueda leer (puede ser un PDF escaneado o de imágenes). "
                "Pide al usuario que copie y pegue el texto en el chat.")
    total = (len(texto) + TAMANO_PARTE - 1) // TAMANO_PARTE
    if parte < 1 or parte > total:
        return f"Solo hay {total} parte(s) en este documento."
    trozo = texto[(parte - 1) * TAMANO_PARTE: parte * TAMANO_PARTE]
    pie = f"\n[Quedan más partes: lee la parte {parte + 1} de {total}]" if parte < total else "\n[Fin del documento]"
    return f"[{ruta.name} · parte {parte} de {total}]\n{trozo}{pie}"


# ---------- Herramientas: edición de fotos ----------
# Las fotos originales están en la carpeta 'fotos'. Las herramientas NUNCA las modifican:
# guardan el resultado como copia en 'fotos/editadas'.

CARPETA_FOTOS = Path("fotos")
CARPETA_EDITADAS = CARPETA_FOTOS / "editadas"
CARPETA_EDITADAS.mkdir(parents=True, exist_ok=True)
EXTENSIONES_FOTOS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"}
MAX_LADO = 8000  # píxeles máximos por lado al redimensionar

SIN_PILLOW = "ERROR: falta la librería Pillow (pip install pillow)."


def _remuestreo():
    return getattr(Image, "Resampling", Image).LANCZOS


def _localizar_foto(nombre: str):
    """Busca la foto primero entre las editadas y luego entre las originales.
    Devuelve (ruta, es_copia_editada) o (None, False)."""
    limpio = Path(nombre).name
    for carpeta, es_copia in ((CARPETA_EDITADAS, True), (CARPETA_FOTOS, False)):
        ruta = carpeta / limpio
        if ruta.is_file() and ruta.suffix.lower() in EXTENSIONES_FOTOS:
            return ruta, es_copia
    return None, False


def _abrir_foto(nombre: str):
    """Abre la foto corrigiendo la rotación de las fotos hechas con el móvil."""
    ruta, es_copia = _localizar_foto(nombre)
    if ruta is None:
        return None, None, False
    imagen = Image.open(ruta)
    imagen = ImageOps.exif_transpose(imagen)
    imagen.load()
    return imagen, ruta, es_copia


def _guardar_foto(imagen, ruta_origen: Path, es_copia: bool, extension: str = "") -> str:
    """Guarda el resultado. Si se editaba una copia, se actualiza esa copia; si era la
    original, se crea una copia nueva (nombre_editada.jpg, nombre_editada_2.jpg...)."""
    extension = (extension or ruta_origen.suffix).lower()
    if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
        extension = ".png"  # bmp, gif... se guardan como png
    if es_copia and extension == ruta_origen.suffix.lower():
        destino = ruta_origen
    else:
        base = ruta_origen.stem.replace("_editada", "") if es_copia else ruta_origen.stem
        destino = CARPETA_EDITADAS / f"{base}_editada{extension}"
        contador = 2
        while destino.exists():
            destino = CARPETA_EDITADAS / f"{base}_editada_{contador}{extension}"
            contador += 1
    if extension in {".jpg", ".jpeg"} and imagen.mode not in ("RGB", "L"):
        fondo = Image.new("RGB", imagen.size, "white")  # el jpg no admite transparencia
        fondo.paste(imagen.convert("RGBA"), mask=imagen.convert("RGBA").split()[-1])
        imagen = fondo
    opciones = {"quality": 90} if extension in {".jpg", ".jpeg", ".webp"} else {}
    imagen.save(destino, **opciones)
    return destino.name


def _resultado(nombre_guardado: str, imagen, detalle: str) -> str:
    return (f"Hecho: {detalle}. Guardada la copia «{nombre_guardado}» ({imagen.width}x{imagen.height}). "
            f"La original no se ha tocado. Para editarla más, usa el nombre «{nombre_guardado}». "
            f"Escribe en tu respuesta, en una línea aparte, [[foto:{nombre_guardado}]] para mostrarla.")


@tool
def listar_fotos() -> str:
    """Lista las fotos subidas por el usuario (originales) y las copias ya editadas,
    con su tamaño en píxeles. Úsala cuando mencione una foto sin dar el nombre exacto."""
    if Image is None:
        return SIN_PILLOW
    partes = []
    for titulo, carpeta in (("Originales", CARPETA_FOTOS), ("Editadas", CARPETA_EDITADAS)):
        lineas = []
        for f in sorted(carpeta.iterdir()):
            if f.is_file() and f.suffix.lower() in EXTENSIONES_FOTOS:
                try:
                    with Image.open(f) as im:
                        lineas.append(f"- {f.name} ({im.width}x{im.height})")
                except Exception:
                    lineas.append(f"- {f.name} (no se puede abrir)")
        partes.append(f"{titulo}:\n" + ("\n".join(lineas) if lineas else "  (ninguna)"))
    return "\n".join(partes)


@tool
def redimensionar_foto(nombre: str, ancho: int = 0, alto: int = 0) -> str:
    """Cambia el tamaño de una foto, en píxeles. Si solo das ancho o solo alto, se mantiene
    la proporción. Si das los dos, la foto se ajusta a ese tamaño exacto (puede deformarse)."""
    if Image is None:
        return SIN_PILLOW
    if ancho <= 0 and alto <= 0:
        return "Indica el ancho, el alto o ambos (en píxeles)."
    if max(ancho, alto) > MAX_LADO:
        return f"El tamaño máximo por lado es {MAX_LADO} píxeles."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    if ancho <= 0:
        ancho = max(1, round(imagen.width * alto / imagen.height))
    if alto <= 0:
        alto = max(1, round(imagen.height * ancho / imagen.width))
    nueva = imagen.resize((ancho, alto), _remuestreo())
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, f"redimensionada a {ancho}x{alto}")


@tool
def recortar_foto(nombre: str, izquierda_pct: float = 0, arriba_pct: float = 0,
                  derecha_pct: float = 0, abajo_pct: float = 0) -> str:
    """Recorta una foto QUITANDO un porcentaje (0-45) de cada lado. Ejemplo: quitar el 10% de
    arriba y de abajo = arriba_pct=10, abajo_pct=10. Para recortar a un cuadrado, calcula los
    porcentajes a partir del tamaño que da listar_fotos."""
    if Image is None:
        return SIN_PILLOW
    valores = [izquierda_pct, arriba_pct, derecha_pct, abajo_pct]
    if any(v < 0 or v > 45 for v in valores) or izquierda_pct + derecha_pct >= 90 or arriba_pct + abajo_pct >= 90:
        return "Los porcentajes deben estar entre 0 y 45 por lado."
    if not any(valores):
        return "Indica cuánto quitar de al menos un lado."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    w, h = imagen.size
    caja = (round(w * izquierda_pct / 100), round(h * arriba_pct / 100),
            w - round(w * derecha_pct / 100), h - round(h * abajo_pct / 100))
    nueva = imagen.crop(caja)
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, "recortada")


@tool
def girar_foto(nombre: str, grados: int = 0, voltear: str = "") -> str:
    """Gira una foto 'grados' en el sentido de las agujas del reloj (90 = a la derecha,
    -90 = a la izquierda, 180 = del revés). 'voltear' puede ser 'horizontal' (espejo) o
    'vertical', o dejarse vacío."""
    if Image is None:
        return SIN_PILLOW
    voltear = voltear.strip().lower()
    if grados % 360 == 0 and voltear not in ("horizontal", "vertical"):
        return "Indica los grados a girar (90, -90, 180...) o voltear='horizontal' / 'vertical'."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    nueva = imagen
    detalle = []
    if grados % 360:
        nueva = nueva.rotate(-grados, expand=True, resample=Image.BICUBIC)
        detalle.append(f"girada {grados}°")
    if voltear == "horizontal":
        nueva = ImageOps.mirror(nueva)
        detalle.append("volteada en horizontal")
    elif voltear == "vertical":
        nueva = ImageOps.flip(nueva)
        detalle.append("volteada en vertical")
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, ", ".join(detalle))


@tool
def ajustar_foto(nombre: str, brillo: float = 1.0, contraste: float = 1.0,
                 saturacion: float = 1.0, nitidez: float = 1.0) -> str:
    """Ajusta una foto con multiplicadores entre 0 y 3, donde 1.0 = sin cambio.
    Ejemplos: más luminosa = brillo 1.3; más colorida = saturacion 1.4; menos color = saturacion 0.6;
    más definida = nitidez 2.0. Cambia solo los valores que haga falta."""
    if Image is None:
        return SIN_PILLOW
    valores = {"brillo": brillo, "contraste": contraste, "saturacion": saturacion, "nitidez": nitidez}
    if any(v < 0 or v > 3 for v in valores.values()):
        return "Los valores deben estar entre 0 y 3 (1.0 = sin cambio)."
    if all(v == 1.0 for v in valores.values()):
        return "No has indicado ningún cambio (todos los valores son 1.0)."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    nueva = imagen.convert("RGBA") if imagen.mode in ("RGBA", "LA", "P") else imagen.convert("RGB")
    for clase, valor in ((ImageEnhance.Brightness, brillo), (ImageEnhance.Contrast, contraste),
                         (ImageEnhance.Color, saturacion), (ImageEnhance.Sharpness, nitidez)):
        if valor != 1.0:
            nueva = clase(nueva).enhance(valor)
    cambios = ", ".join(f"{k} {v}" for k, v in valores.items() if v != 1.0)
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, f"ajustada ({cambios})")


@tool
def filtro_foto(nombre: str, filtro: str) -> str:
    """Aplica un filtro a una foto. Filtros: 'blanco_negro', 'sepia', 'invertir',
    'desenfocar' (suaviza) o 'enfocar' (más nitidez)."""
    if Image is None:
        return SIN_PILLOW
    filtro = filtro.strip().lower().replace(" ", "_")
    filtros_validos = {"blanco_negro", "sepia", "invertir", "desenfocar", "enfocar"}
    if filtro not in filtros_validos:
        return "Filtro no válido. Usa: " + ", ".join(sorted(filtros_validos)) + "."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    base = imagen.convert("RGB")
    if filtro == "blanco_negro":
        nueva = ImageOps.grayscale(base).convert("RGB")
    elif filtro == "sepia":
        nueva = ImageOps.colorize(ImageOps.grayscale(base), black="#201005", white="#ffe9c8", mid="#9b7653")
    elif filtro == "invertir":
        nueva = ImageOps.invert(base)
    elif filtro == "desenfocar":
        nueva = base.filter(ImageFilter.GaussianBlur(radius=max(2, min(base.size) / 120)))
    else:
        nueva = base.filter(ImageFilter.UnsharpMask(radius=2, percent=160, threshold=3))
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, f"filtro «{filtro}» aplicado")


def _fuente(tamano: int):
    for nombre in ("arial.ttf", "DejaVuSans.ttf", "Arial.ttf"):
        try:
            return ImageFont.truetype(nombre, tamano)
        except OSError:
            continue
    try:
        return ImageFont.load_default(tamano)
    except TypeError:
        return ImageFont.load_default()


@tool
def anadir_texto_foto(nombre: str, texto: str, posicion: str = "abajo",
                      color: str = "white", tamano_pct: float = 6) -> str:
    """Escribe un texto sobre una foto. posicion: 'arriba', 'centro' o 'abajo'.
    color: nombre en inglés ('white', 'black', 'red', 'yellow'...) o hexadecimal ('#ff8800').
    tamano_pct: altura de la letra como porcentaje de la altura de la foto (2-20)."""
    if Image is None:
        return SIN_PILLOW
    if not texto.strip() or len(texto) > 200:
        return "Escribe un texto (máximo 200 caracteres)."
    posicion = posicion.strip().lower()
    if posicion not in ("arriba", "centro", "abajo"):
        return "La posición debe ser 'arriba', 'centro' o 'abajo'."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    try:
        ImageColor.getrgb(color)
    except ValueError:
        color = "white"
    nueva = imagen.convert("RGB")
    dibujo = ImageDraw.Draw(nueva)
    tamano = max(12, int(nueva.height * min(max(tamano_pct, 2), 20) / 100))
    fuente = _fuente(tamano)
    # Si el texto no cabe a lo ancho, se reduce la letra
    while tamano > 12 and dibujo.textbbox((0, 0), texto, font=fuente)[2] > nueva.width * 0.94:
        tamano = int(tamano * 0.9)
        fuente = _fuente(tamano)
    caja = dibujo.textbbox((0, 0), texto, font=fuente)
    ancho_texto, alto_texto = caja[2] - caja[0], caja[3] - caja[1]
    x = (nueva.width - ancho_texto) / 2 - caja[0]
    margen = nueva.height * 0.04
    y = {"arriba": margen, "centro": (nueva.height - alto_texto) / 2, "abajo": nueva.height - alto_texto - margen * 1.5}[posicion] - caja[1]
    dibujo.text((x, y), texto, font=fuente, fill=color,
                stroke_width=max(1, tamano // 14), stroke_fill="black")
    return _resultado(_guardar_foto(nueva, ruta, es_copia), nueva, f"texto «{texto[:30]}» añadido")


@tool
def convertir_foto(nombre: str, formato: str, calidad: int = 85) -> str:
    """Convierte una foto a 'jpg', 'png' o 'webp'. Con 'jpg' o 'webp', 'calidad' (30-95) controla
    el peso: más baja = archivo más pequeño. Úsala también para comprimir/reducir el peso de una foto."""
    if Image is None:
        return SIN_PILLOW
    formato = formato.strip().lower().lstrip(".")
    if formato not in ("jpg", "jpeg", "png", "webp"):
        return "Formato no válido. Usa 'jpg', 'png' o 'webp'."
    imagen, ruta, es_copia = _abrir_foto(nombre)
    if imagen is None:
        return f"No encuentro la foto '{nombre}'."
    extension = ".jpg" if formato == "jpeg" else f".{formato}"
    calidad = min(max(int(calidad), 30), 95)
    # Se guarda a mano para poder usar la calidad pedida
    base = ruta.stem.replace("_editada", "") if es_copia else ruta.stem
    destino = CARPETA_EDITADAS / f"{base}_editada{extension}"
    contador = 2
    while destino.exists():
        destino = CARPETA_EDITADAS / f"{base}_editada_{contador}{extension}"
        contador += 1
    salida = imagen
    if extension == ".jpg" and imagen.mode not in ("RGB", "L"):
        fondo = Image.new("RGB", imagen.size, "white")
        rgba = imagen.convert("RGBA")
        fondo.paste(rgba, mask=rgba.split()[-1])
        salida = fondo
    opciones = {"quality": calidad} if extension in (".jpg", ".webp") else {}
    salida.save(destino, **opciones)
    peso_antes, peso_despues = ruta.stat().st_size // 1024, destino.stat().st_size // 1024
    return (f"Hecho: convertida a {formato.upper()} (calidad {calidad}). Guardada la copia «{destino.name}» "
            f"({peso_antes} KB → {peso_despues} KB). La original no se ha tocado. "
            f"Escribe en tu respuesta, en una línea aparte, [[foto:{destino.name}]] para mostrarla.")


# ---------- El agente ----------

INSTRUCCIONES = """Eres un asistente personal de organización y estudio. Hablas en español,
de forma cercana y clara. Eres breve en las gestiones sencillas (guardar una tarea, un saludo) y
más completo cuando tienes que explicar algo.

Reglas:
- Antes de hablar de fechas ("hoy", "mañana", "esta semana") consulta fecha_hora_actual.
- Cuando el usuario mencione algo que tiene que hacer, ofrécete a guardarlo como tarea.
- Para el "resumen del día": consulta la fecha, lista las tareas pendientes y propón un
  orden de prioridades realista (máximo 3 cosas importantes).
- Para estudiar: si el usuario pega apuntes, guárdalos; si pide tarjetas, genéralas desde los
  apuntes y guárdalas; si pide repasar, hazle UNA pregunta cada vez y corrige su respuesta.
- No inventes tareas ni apuntes que no existan: consulta siempre las herramientas.
- Cuando el usuario pida ver sus apuntes, tareas o tarjetas, MUESTRA en tu respuesta el
  contenido completo que devuelva la herramienta (con sus números), copiado tal cual.
  Nunca digas solo "aquí tienes" sin escribir el contenido. Si la herramienta dice que
  no hay nada, díselo claramente.
- Solo puedes eliminar apuntes (con eliminar_apunte), nunca tareas ni tarjetas. Antes de
  eliminar, usa leer_apuntes para identificar el apunte correcto y bórralo de uno en uno.
  El programa pedirá confirmación al usuario; no borres si el usuario no lo ha pedido claramente.
- Si te piden borrar tareas o tarjetas, explica que debe hacerlo a mano.
- Revisión de código: si el usuario pide revisar o corregir código, usa listar_archivos_codigo
  y leer_archivo_codigo (también puede pegarte trozos cortos en el chat). Responde así:
  1) di en una frase si funciona o no; 2) lista los errores indicando el número de línea y
  explicándolos con palabras sencillas; 3) muestra la versión corregida de las partes
  afectadas; 4) añade como máximo 2 mejoras opcionales. Si el usuario
  pide CORREGIR o ARREGLAR el archivo, después de leerlo usa guardar_codigo_corregido con el
  archivo completo corregido (sin números de línea y sin dejar partes fuera). El programa
  enseñará los cambios y pedirá confirmación. Nunca modifiques un archivo si solo te pidieron
  revisarlo. Si algo no está claro, pregunta qué debía hacer el código.
- Documentos: el usuario puede subir archivos (PDF, Word, texto) a la carpeta 'documentos'. Si
  menciona un archivo, un documento o un PDF, o dice que ha subido algo, usa listar_documentos y
  leer_documento (si es largo, lee TODAS las partes antes de responder). Después responde así:
  1) resumen de qué trata en 2-3 frases; 2) QUÉ TIENE QUE HACER el usuario, en pasos numerados,
  con fechas y requisitos si aparecen; 3) ofrécete a guardar esos pasos como tareas y a ayudarle
  con cada uno. Basa todo SOLO en el contenido del documento: si algo no aparece, dilo en lugar
  de inventarlo.
- Cuando el usuario te pida ayuda con un paso o una actividad, ayúdale de forma práctica:
  redacta borradores, explica conceptos con ejemplos, hazle preguntas para comprobar que lo
  entiende. Si hay varias formas de hacerlo, pregunta cuál prefiere.
- Fotos: el usuario puede subir fotos (JPG, PNG, WEBP...) a la carpeta 'fotos'. Para editarlas usa
  listar_fotos y las herramientas redimensionar_foto, recortar_foto, girar_foto, ajustar_foto,
  filtro_foto, anadir_texto_foto y convertir_foto. Nunca modifican la original: crean una copia en
  'fotos/editadas'. Para encadenar varios cambios, usa el nombre de la copia que te devolvió la
  herramienta. Tú NO puedes ver la foto, solo conoces su tamaño: no describas su contenido ni
  inventes qué aparece en ella. Traduce lo que pide el usuario a valores razonables (por ejemplo,
  "más luminosa" = brillo 1.3) y dile qué valores has usado. Cuando termines de editar, o si el
  usuario quiere ver una foto, escribe en tu respuesta, en una línea aparte, [[foto:NOMBRE]] con el
  nombre exacto del archivo: así la aplicación se la enseña. Si la petición no se puede hacer con
  esas herramientas (quitar el fondo, retocar caras...), dilo claramente.
- Cómo explicar (cuando el usuario te pida explicar, enseñar, resumir o aclarar algo):
  1) Empieza por la idea principal en una o dos frases sencillas, antes de entrar en detalles.
  2) Usa palabras normales. Si necesitas un término técnico, defínelo la primera vez.
  3) Incluye siempre un ejemplo concreto de la vida real o una analogía.
  4) Si hay pasos o partes, ordénalos y numéralos. Una sola idea nueva por párrafo, de lo simple
     a lo complejo; nunca un bloque enorme de texto.
  5) Adapta el nivel al usuario. Si no sabes cuánto sabe del tema, empieza sencillo y pregúntale si
     quiere más detalle. Si dice que no lo entiende, NO repitas lo mismo con otras palabras:
     explícalo de otra forma, con otro ejemplo y más despacio.
  6) En temas de estudio, termina con "Lo importante:" en 2-3 líneas y ofrécete a hacerle 2-3
     preguntas para comprobar que lo ha entendido, o a crear tarjetas de repaso.
  7) No inventes datos, cifras, fechas ni citas. Si no estás seguro de algo, dilo. En temas de salud,
     leyes o dinero explica lo general y recuérdale que lo consulte con un profesional.
  8) Si explicas a partir de un documento, apuntes o código, básate SOLO en lo que contienen e
     indica de dónde sale cada cosa (página, apartado o número de línea).
  9) Longitud justa: corto para preguntas simples; más desarrollado solo si el tema lo pide.
- Formato: escribe en texto plano. No uses Markdown (nada de **negritas**, # títulos ni tablas).
  Para estructurar usa frases cortas, listas con guiones o con números y líneas en blanco entre
  secciones. El código sí va dentro de un bloque que empieza y termina con tres comillas
  invertidas (```)."""

# Si te da error de modelo, prueba con "gemini-3.5-flash"
modelo = ChatGoogleGenerativeAI(model="gemini-3.5-flash-lite")

agente = create_react_agent(
    modelo,
    [
        fecha_hora_actual,
        anadir_tarea,
        listar_tareas,
        completar_tarea,
        guardar_apunte,
        leer_apuntes,
        eliminar_apunte,
        guardar_tarjetas,
        leer_tarjetas,
        listar_archivos_codigo,
        leer_archivo_codigo,
        guardar_codigo_corregido,
        listar_documentos,
        leer_documento,
        listar_fotos,
        redimensionar_foto,
        recortar_foto,
        girar_foto,
        ajustar_foto,
        filtro_foto,
        anadir_texto_foto,
        convertir_foto,
    ],
    prompt=INSTRUCCIONES,  # en versiones antiguas de langgraph se llama state_modifier
)


def a_texto(contenido) -> str:
    """Gemini a veces devuelve la respuesta como lista de bloques; esto la convierte en texto."""
    if isinstance(contenido, str):
        return contenido
    partes = []
    for bloque in contenido:
        if isinstance(bloque, str):
            partes.append(bloque)
        elif isinstance(bloque, dict) and bloque.get("type") == "text":
            partes.append(bloque.get("text", ""))
    return "".join(partes)


def main() -> None:
    print("Agente personal (Gemini) listo. Escribe 'salir' para terminar.\n")
    historial = []
    while True:
        entrada = input("Tú: ").strip()
        if entrada.lower() in {"salir", "exit", "quit"}:
            break
        if not entrada:
            continue
        historial.append(("user", entrada))
        try:
            resultado = agente.invoke(
                {"messages": historial},
                config={"recursion_limit": 15},  # límite de pasos por seguridad
            )
        except Exception as error:
            print(f"\nError: {error}\n(Si menciona '429' o 'quota', has llegado al límite gratuito: espera un rato.)\n")
            historial.pop()
            continue
        historial = resultado["messages"]
        print(f"\nAgente: {a_texto(historial[-1].content)}\n")


if __name__ == "__main__":
    main()
